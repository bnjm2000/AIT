"""Asset template import: parsing, signed previews, validation, and atomic saves.

Shared inventory ID, audit, department, and locking rules come from the app.
The supplied data manager stays a request-local proxy.
"""

import copy
import csv
import hashlib
import io
import json
import os
import re
from datetime import datetime

from flask import jsonify, request, send_file, session
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

from models import InventoryItem, normalize_asset_tags

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None

ASSET_IMPORT_MAX_ROWS = 1000
ASSET_IMPORT_MAX_BYTES = 5 * 1024 * 1024
ASSET_IMPORT_MAX_RECORDS = max(
    ASSET_IMPORT_MAX_ROWS,
    int(os.environ.get('ASSET_IMPORT_MAX_RECORDS', '50000')),
)
ASSET_IMPORT_PLAN_MAX_AGE_SECONDS = 30 * 60
ASSET_IMPORT_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    'outputs',
    'asset-import',
    'asset-import-template.csv',
)
ASSET_IMPORT_FILE_EXTENSIONS = {'.csv', '.xlsx', '.xlsm', '.xls'}
ASSET_IMPORT_HEADERS = {
    'brand': ('brand',),
    'model': ('model', 'model number'),
    'description': ('description',),
    'version': ('version',),
    'department': (
        'department',
        'department code',
        'department name',
        'department code or name',
        'dept',
        'dept code',
        'dept name',
    ),
    'quantity': ('quantity', 'qty'),
    'isBulk': ('bulk asset', 'bulk', 'is bulk'),
    'serials': ('primary serial numbers', 'primary serials', 'serial numbers', 'serials'),
    'secondarySerials': ('secondary serial numbers', 'secondary serials', 'second serial numbers'),
    'dateOfPurchase': ('date of purchase', 'purchase date'),
    'notes': ('notes',),
    'tags': ('tags',),
    'defaultLocation': ('default location', 'location'),
    'assetIdPrefix': ('custom asset id prefix', 'asset id prefix', 'custom prefix'),
}


def register_asset_import_routes(
    app, *, data_manager, require_admin, logger, invalidate_cache, log_action,
    mark_data_snapshot_current, mark_realtime_change, asset_audit_timestamp, asset_audit_user,
    duplicate_serial_candidates, duplicate_serial_details, build_duplicate_serial_index,
    duplicate_serial_warning, asset_id_plan, serial_count_mismatches, current_company_code,
    department_payload, department_record, inventory_action_lock, is_bulk_asset,
    load_departments, mark_asset_created, normalise_asset_id_prefix,
    normalise_asset_purchase_date, normalise_department_code, safe_int, save_departments,
):
    def _normalise_asset_import_header(value):
        return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', str(value or '').casefold())).strip()


    def _asset_import_header_map(values):
        aliases = {
            _normalise_asset_import_header(alias): field
            for field, field_aliases in ASSET_IMPORT_HEADERS.items()
            for alias in field_aliases
        }
        return {
            aliases[normalized]: index
            for index, value in enumerate(values or [])
            if (normalized := _normalise_asset_import_header(value)) in aliases
        }


    def _asset_import_serials(value):
        if value is None:
            return []
        if isinstance(value, (list, tuple, set)):
            serials = []
            for item in value:
                if item is None or (isinstance(item, str) and not item.strip()):
                    serials.append('')
                else:
                    serials.extend(_asset_import_serials(item))
            while serials and not serials[-1]:
                serials.pop()
            return serials
        if isinstance(value, bool):
            return [str(value)]
        if isinstance(value, int):
            return [str(value)]
        if isinstance(value, float):
            return [str(int(value)) if value.is_integer() else format(value, 'g')]
        text = str(value).replace('\r\n', '\n').replace('\r', '\n')
        serials = [item.strip() for item in re.split(r'[,;\n]', text)]
        while serials and not serials[-1]:
            serials.pop()
        return serials


    def _asset_import_bool(value):
        normalized = str(value or '').strip().casefold()
        if normalized in ('', 'no', 'n', 'false', '0'):
            return False
        if normalized in ('yes', 'y', 'true', '1'):
            return True
        raise ValueError('Bulk Asset must be Yes or No')


    def _asset_import_quantity(value):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError('Quantity must be a whole number')
        quantity = safe_int(value, 0)
        if quantity < 1 or quantity > 500:
            raise ValueError('Quantity must be between 1 and 500')
        return quantity


    def _asset_import_unchecked_row_payload(values, header_map):
        def read(field):
            index = header_map.get(field)
            return values[index] if index is not None and index < len(values) else None

        return {
            'brand': str(read('brand') or '').strip(),
            'model': str(read('model') or '').strip(),
            'description': str(read('description') or '').strip(),
            'version': str(read('version') or '').strip(),
            'department': str(read('department') or '').strip(),
            'quantity': read('quantity'),
            'isBulk': read('isBulk'),
            'serials': _asset_import_serials(read('serials')),
            'secondarySerials': _asset_import_serials(read('secondarySerials')),
            'dateOfPurchase': read('dateOfPurchase').strftime('%Y-%m-%d') if hasattr(read('dateOfPurchase'), 'strftime') else str(read('dateOfPurchase') or '').strip(),
            'notes': str(read('notes') or '').strip(),
            'tags': normalize_asset_tags(read('tags')),
            'defaultLocation': str(read('defaultLocation') or '').strip() or 'Store',
            'assetIdPrefix': str(read('assetIdPrefix') or '').strip(),
            'createDepartment': False,
            'newDepartmentCode': '',
            'newDepartmentName': '',
        }


    def _parse_asset_import_rows(rows):
        rows = iter(rows)
        buffered_rows = []
        for row_number, values in enumerate(rows, start=1):
            buffered_rows.append((row_number, list(values or [])))
            if row_number >= 12:
                break

        header_row_number = None
        header_map = {}
        for row_number, values in buffered_rows:
            candidate = _asset_import_header_map(values)
            if 'brand' in candidate and 'model' in candidate:
                header_row_number = row_number
                header_map = candidate
                break
        if header_row_number is None:
            raise ValueError('Could not find the asset column headers in the uploaded file')

        missing_headers = [
            label for field, label in (
                ('brand', 'Brand'),
                ('model', 'Model'),
                ('department', 'Department Code or Name'),
                ('quantity', 'Quantity'),
            ) if field not in header_map
        ]
        if missing_headers:
            raise ValueError(f"Missing required column(s): {', '.join(missing_headers)}")

        data_rows = (
            [(number, values) for number, values in buffered_rows if number > header_row_number]
            + [(number, list(values or [])) for number, values in enumerate(rows, start=len(buffered_rows) + 1)]
        )
        parsed_rows = []
        rejected = []
        for row_number, values in data_rows:
            if not any(value not in (None, '') for value in values):
                continue
            if len(parsed_rows) >= ASSET_IMPORT_MAX_ROWS:
                raise ValueError(f'An import can contain at most {ASSET_IMPORT_MAX_ROWS} rows')
            try:
                payload = _normalise_asset_import_payload(
                    _asset_import_unchecked_row_payload(values, header_map),
                    allow_unmatched_department=True,
                )
                payload['sourceRow'] = row_number
                parsed_rows.append(payload)
            except ValueError as error:
                payload = _asset_import_unchecked_row_payload(values, header_map)
                payload['sourceRow'] = row_number
                payload['validationError'] = str(error)
                parsed_rows.append(payload)
                rejected.append({'sourceRow': row_number, 'error': str(error)})

        if not parsed_rows:
            raise ValueError('The uploaded file does not contain any asset rows')
        return {'rows': parsed_rows, 'rejected': rejected}


    def _parse_asset_import_workbook(file_bytes):
        if load_workbook is None:
            raise RuntimeError('Excel import support is not installed')
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as error:
            raise ValueError('The uploaded file is not a valid .xlsx workbook') from error

        worksheet = workbook['Assets'] if 'Assets' in workbook.sheetnames else workbook.active
        return _parse_asset_import_rows(worksheet.iter_rows(values_only=True))


    def _decode_asset_import_csv(file_bytes):
        if file_bytes.startswith((b'\xff\xfe', b'\xfe\xff')):
            try:
                return file_bytes.decode('utf-16')
            except UnicodeDecodeError as error:
                raise ValueError('The CSV file contains invalid UTF-16 text') from error

        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError('The CSV file must use UTF-8, UTF-16, or Windows-1252 text encoding')


    def _parse_asset_import_legacy_excel(file_bytes):
        if xlrd is None:
            raise RuntimeError('Legacy .xls import support is not installed')
        try:
            workbook = xlrd.open_workbook(file_contents=file_bytes)
            worksheet = workbook.sheet_by_name('Assets') if 'Assets' in workbook.sheet_names() else workbook.sheet_by_index(0)
        except Exception as error:
            raise ValueError('The uploaded file is not a valid .xls workbook') from error

        def rows():
            for row_index in range(worksheet.nrows):
                values = []
                for cell in worksheet.row(row_index):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                    else:
                        values.append(cell.value)
                yield values

        return _parse_asset_import_rows(rows())


    def _parse_asset_import_file(file_bytes, filename):
        extension = os.path.splitext(str(filename or ''))[1].lower()
        if extension == '.csv':
            text = _decode_asset_import_csv(file_bytes)
            return _parse_asset_import_rows(csv.reader(io.StringIO(text, newline='')))
        if extension in ('.xlsx', '.xlsm'):
            return _parse_asset_import_workbook(file_bytes)
        if extension == '.xls':
            return _parse_asset_import_legacy_excel(file_bytes)
        raise ValueError('Upload a CSV or Excel file (.csv, .xlsx, .xlsm, or .xls)')


    def _asset_import_department_resolution(value, departments=None):
        departments = load_departments() if departments is None else departments
        raw_value = str(value or '').strip()
        if not raw_value:
            return {
                'input': '',
                'matched': False,
                'ambiguous': False,
                'code': '',
                'name': '',
                'suggestedCode': '',
                'issue': 'Department is required',
            }

        normalized_code = normalise_department_code(raw_value)
        if normalized_code in departments:
            department = departments[normalized_code]
            return {
                'input': raw_value,
                'matched': True,
                'ambiguous': False,
                'code': normalized_code,
                'name': str(department.get('name') or normalized_code).strip() or normalized_code,
                'suggestedCode': normalized_code,
                'issue': '',
            }

        name_matches = [
            (code, department)
            for code, department in departments.items()
            if str(department.get('name') or code).strip().casefold() == raw_value.casefold()
        ]
        if len(name_matches) == 1:
            code, department = name_matches[0]
            return {
                'input': raw_value,
                'matched': True,
                'ambiguous': False,
                'code': code,
                'name': str(department.get('name') or code).strip() or code,
                'suggestedCode': code,
                'issue': '',
            }

        if len(name_matches) > 1:
            matching_codes = ', '.join(sorted(code for code, _ in name_matches))
            return {
                'input': raw_value,
                'matched': False,
                'ambiguous': True,
                'code': '',
                'name': '',
                'suggestedCode': '',
                'issue': f'Department name "{raw_value}" matches more than one code ({matching_codes}); enter a department code instead',
            }

        return {
            'input': raw_value,
            'matched': False,
            'ambiguous': False,
            'code': '',
            'name': '',
            'suggestedCode': normalized_code,
            'issue': f'Department "{raw_value}" does not match an existing department',
        }


    def _normalise_asset_import_payload(data, allow_unmatched_department=False, departments=None):
        data = data or {}
        departments = load_departments() if departments is None else departments
        department_resolution = _asset_import_department_resolution(data.get('department'), departments)
        create_department = _asset_import_bool(data.get('createDepartment'))
        new_department_code = normalise_department_code(
            data.get('newDepartmentCode') or department_resolution.get('suggestedCode')
        )
        new_department_name = str(
            data.get('newDepartmentName') or department_resolution.get('input')
        ).strip()

        if department_resolution['matched']:
            department_code = department_resolution['code']
            create_department = False
            new_department_code = ''
            new_department_name = ''
        elif create_department:
            if department_resolution['ambiguous']:
                raise ValueError(department_resolution['issue'])
            if not new_department_code:
                raise ValueError('New department code is required')
            if not new_department_name:
                raise ValueError('New department name is required')
            if new_department_code in departments:
                existing_name = departments[new_department_code].get('name') or new_department_code
                raise ValueError(
                    f'Department code {new_department_code} already exists as {existing_name}; use the existing code or name instead'
                )
            existing_name_match = next((
                code for code, department in departments.items()
                if str(department.get('name') or code).strip().casefold() == new_department_name.casefold()
            ), '')
            if existing_name_match:
                raise ValueError(
                    f'Department name {new_department_name} already exists as {existing_name_match}; use the existing code or name instead'
                )
            department_code = new_department_code
        elif allow_unmatched_department:
            department_code = department_resolution['input']
        else:
            raise ValueError(department_resolution['issue'])

        payload = {
            'brand': str(data.get('brand') or '').strip(),
            'model': str(data.get('model') or '').strip(),
            'description': str(data.get('description') or '').strip(),
            'version': str(data.get('version') or '').strip(),
            'department': department_code,
            'departmentInput': str(data.get('departmentInput') or department_resolution['input']).strip(),
            'departmentMatched': department_resolution['matched'],
            'departmentMatchedName': department_resolution['name'],
            'departmentIssue': '' if department_resolution['matched'] or create_department else department_resolution['issue'],
            'createDepartment': create_department,
            'newDepartmentCode': new_department_code,
            'newDepartmentName': new_department_name,
            'quantity': _asset_import_quantity(data.get('quantity')),
            'isBulk': _asset_import_bool(data.get('isBulk')),
            'serials': _asset_import_serials(data.get('serials')),
            'secondarySerials': _asset_import_serials(data.get('secondarySerials')),
            'dateOfPurchase': normalise_asset_purchase_date(data.get('dateOfPurchase')),
            'notes': str(data.get('notes') or '').strip(),
            'tags': normalize_asset_tags(data.get('tags')),
            'defaultLocation': str(data.get('defaultLocation') or '').strip() or 'Store',
            'assetIdPrefix': normalise_asset_id_prefix(data.get('assetIdPrefix')),
            'sourceRow': safe_int(data.get('sourceRow'), 0),
        }
        if not payload['brand']:
            raise ValueError('Brand is required')
        if not payload['model']:
            raise ValueError('Model number is required')
        if not payload['department']:
            raise ValueError('Department is required')
        if payload['isBulk']:
            payload['serials'] = []
            payload['secondarySerials'] = []
            payload['assetIdPrefix'] = ''
        return payload


    def _asset_import_validation_warnings(payload):
        warnings = []
        for mismatch in serial_count_mismatches(
            payload,
            payload.get('quantity', 0),
            payload.get('isBulk', False),
        ):
            count = mismatch['count']
            quantity = mismatch['quantity']
            label = mismatch['type'].capitalize()
            if count < quantity:
                missing = quantity - count
                warnings.append(
                    f'{label} serial numbers: {count}/{quantity}; '
                    f'{missing} asset{"" if missing == 1 else "s"} will be saved without one'
                )
            else:
                extra = count - quantity
                warnings.append(
                    f'{label} serial numbers: {count}/{quantity}; '
                    f'{extra} extra serial number{"" if extra == 1 else "s"} will not be saved'
                )

        purchase_date = str(payload.get('dateOfPurchase') or '').strip()
        if purchase_date and datetime.strptime(purchase_date, '%Y-%m-%d').date() > datetime.now().date():
            warnings.append('Date of purchase is in the future')
        return warnings


    def _asset_import_auto_ids(payload, reserved_ids=None):
        reserved_ids = set(reserved_ids or [])
        plan = asset_id_plan(payload)
        if plan['isBulk']:
            used_ids = set(data_manager.inventory.keys()) | reserved_ids
            used_numbers = [
                safe_int(str(asset_id).replace('BULK-', ''), 0)
                for asset_id in used_ids
                if str(asset_id).startswith('BULK-')
            ]
            return [f"BULK-{(max(used_numbers, default=0) + 1):04d}"]

        prefix = plan['prefix']
        pattern = re.compile(rf'^{re.escape(prefix)}#(\d+)$', re.IGNORECASE)
        used_numbers = []
        widths = [2]
        for asset_id in set(data_manager.inventory.keys()) | reserved_ids:
            match = pattern.match(str(asset_id))
            if not match:
                continue
            used_numbers.append(safe_int(match.group(1), 0))
            widths.append(len(match.group(1)))
        start = max(used_numbers, default=0) + 1
        width = max(*widths, len(str(start + payload['quantity'] - 1)))
        return [
            f"{prefix}#{number:0{width}d}"
            for number in range(start, start + payload['quantity'])
        ]


    def _asset_import_different_descriptions(payload):
        brand = str(payload.get('brand') or '').strip().casefold()
        model = str(payload.get('model') or '').strip().casefold()
        description = str(payload.get('description') or '').strip()
        matches = {
            str(getattr(asset, 'description', '') or '').strip() or '(blank description)'
            for asset in data_manager.inventory.values()
            if str(getattr(asset, 'brand', '') or '').strip().casefold() == brand
            and str(getattr(asset, 'model_number', '') or '').strip().casefold() == model
            and str(getattr(asset, 'description', '') or '').strip() != description
        }
        return sorted(matches, key=str.casefold)


    def _asset_import_preview_rows(rows):
        departments = load_departments()
        proposed_departments = _asset_import_proposed_departments(rows, departments)
        planning_departments = {**departments, **proposed_departments}
        preview_rows = []
        reserved_ids = set()
        duplicate_serial_index = build_duplicate_serial_index()
        for row in rows:
            try:
                payload = _normalise_asset_import_payload(
                    row,
                    allow_unmatched_department=True,
                    departments=planning_departments,
                )
                department_code = payload.get('department', '')
                payload['departmentWillBeCreated'] = department_code in proposed_departments
                if _asset_import_bool((row or {}).get('createDepartment')):
                    payload['createDepartment'] = department_code in proposed_departments
                    if payload['createDepartment']:
                        payload['newDepartmentCode'] = department_code
                        payload['newDepartmentName'] = proposed_departments[department_code]['name']
                preview_ids = _asset_import_auto_ids(payload, reserved_ids)
            except ValueError as error:
                preview_rows.append({
                    **row,
                    'validationError': str(error),
                    'suggestedAssetIds': [],
                    'assetIdsPreview': [],
                    'idIssue': '',
                    'differentDescriptions': [],
                    'validationWarnings': [],
                })
                continue
            reserved_ids.update(preview_ids)
            primary_serials = (
                payload['serials'] + [''] * payload['quantity']
            )[:payload['quantity']]
            secondary_serials = (
                payload['secondarySerials'] + [''] * payload['quantity']
            )[:payload['quantity']]
            duplicate_serials = duplicate_serial_details([
                candidate
                for index, asset_id in enumerate(preview_ids)
                for candidate in duplicate_serial_candidates(
                    asset_id,
                    payload['brand'],
                    payload['model'],
                    payload['description'],
                    primary_serials[index],
                    secondary_serials[index],
                )
            ], duplicate_serial_index) if not payload['isBulk'] else []
            preview_rows.append({
                **payload,
                'suggestedAssetIds': preview_ids,
                'assetIdsPreview': preview_ids,
                'idIssue': '',
                'differentDescriptions': _asset_import_different_descriptions(payload),
                'duplicateSerials': duplicate_serials,
                'validationWarnings': [
                    *_asset_import_validation_warnings(payload),
                    *(duplicate_serial_warning(detail) for detail in duplicate_serials),
                ],
                'validationError': '',
            })
        return preview_rows


    def _asset_import_proposed_departments(rows, departments=None):
        departments = load_departments() if departments is None else departments
        new_departments = {}
        new_names = {}
        for row in rows:
            row = row if isinstance(row, dict) else {}
            if not _asset_import_bool(row.get('createDepartment')):
                continue
            resolution = _asset_import_department_resolution(row.get('department'), departments)
            if resolution['matched']:
                continue
            if resolution['ambiguous']:
                raise ValueError(resolution['issue'])
            code = normalise_department_code(
                row.get('newDepartmentCode') or resolution.get('suggestedCode')
            )
            name = str(row.get('newDepartmentName') or resolution.get('input')).strip()
            if not code:
                raise ValueError('New department code is required')
            if not name:
                raise ValueError('New department name is required')
            if code in departments:
                existing_name = departments[code].get('name') or code
                raise ValueError(
                    f'Department code {code} already exists as {existing_name}; '
                    'use the existing code or name instead'
                )
            existing_name_match = next((
                existing_code
                for existing_code, department in departments.items()
                if str(department.get('name') or existing_code).strip().casefold() == name.casefold()
            ), '')
            if existing_name_match:
                raise ValueError(
                    f'Department name {name} already exists as {existing_name_match}; '
                    'use the existing code or name instead'
                )
            name_key = name.casefold()
            existing = new_departments.get(code)
            if existing and existing['name'].casefold() != name_key:
                raise ValueError(
                    f'New department code {code} has conflicting names in this import'
                )
            existing_code = new_names.get(name_key)
            if existing_code and existing_code != code:
                raise ValueError(
                    f'New department name {name} has conflicting codes in this import'
                )
            new_departments[code] = department_record(code, name)
            new_names[name_key] = code
        return new_departments


    def _asset_import_inventory_revision():
        inventory_rows = [
            (
                str(asset_id),
                str(getattr(asset, 'brand', '') or ''),
                str(getattr(asset, 'model_number', '') or ''),
                str(getattr(asset, 'description', '') or ''),
                str(getattr(asset, 'serial_number', '') or ''),
                str(getattr(asset, 'secondary_serial_number', '') or ''),
                bool(is_bulk_asset(asset)),
            )
            for asset_id, asset in data_manager.inventory.items()
            if asset is not None
        ]
        encoded = json.dumps(
            sorted(inventory_rows, key=lambda row: row[0].casefold()),
            ensure_ascii=False,
            separators=(',', ':'),
        ).encode('utf-8')
        return hashlib.sha256(encoded).hexdigest()


    def _asset_import_department_revision(departments=None):
        departments = load_departments() if departments is None else departments
        rows = [
            (
                normalise_department_code(code),
                str((department or {}).get('name') or code).strip(),
            )
            for code, department in (departments or {}).items()
        ]
        encoded = json.dumps(
            sorted(rows, key=lambda row: row[0].casefold()),
            ensure_ascii=False,
            separators=(',', ':'),
        ).encode('utf-8')
        return hashlib.sha256(encoded).hexdigest()


    def _asset_import_plan_rows_digest(rows):
        fields = (
            'brand', 'model', 'description', 'version', 'department', 'quantity',
            'isBulk', 'serials', 'secondarySerials', 'dateOfPurchase', 'notes',
            'tags', 'defaultLocation', 'assetIdPrefix', 'sourceRow',
            'createDepartment', 'newDepartmentCode', 'newDepartmentName',
            'assetIdsPreview',
        )
        canonical = [
            {field: row.get(field) for field in fields}
            for row in rows
        ]
        encoded = json.dumps(
            canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(',', ':'),
        ).encode('utf-8')
        return hashlib.sha256(encoded).hexdigest()


    def _asset_import_plan_serializer():
        return URLSafeTimedSerializer(app.secret_key, salt='asset-import-plan-v1')


    def _asset_import_requested_record_count(rows):
        total = 0
        for row in rows:
            row = row if isinstance(row, dict) else {}
            try:
                quantity = _asset_import_quantity(row.get('quantity'))
                is_bulk = _asset_import_bool(row.get('isBulk'))
            except ValueError:
                continue
            total += 1 if is_bulk else quantity
            if total > ASSET_IMPORT_MAX_RECORDS:
                raise ValueError(
                    f'This import would create more than {ASSET_IMPORT_MAX_RECORDS:,} '
                    'inventory records. Split it into smaller imports.'
                )
        return total


    def _asset_import_plan_response(rows):
        record_count = _asset_import_requested_record_count(rows)
        planned_rows = _asset_import_preview_rows(rows)
        revision = _asset_import_inventory_revision()
        department_revision = _asset_import_department_revision()
        token = _asset_import_plan_serializer().dumps({
            'company': current_company_code(),
            'user': str(session.get('user') or ''),
            'revision': revision,
            'departmentRevision': department_revision,
            'rowsDigest': _asset_import_plan_rows_digest(planned_rows),
        })
        return {
            'rows': planned_rows,
            'planToken': token,
            'inventoryRevision': revision,
            'departmentRevision': department_revision,
            'inventoryRecordsPlanned': record_count,
        }


    def _verify_asset_import_plan(rows, token):
        if not token:
            raise ValueError('Refresh the Asset ID preview before importing')
        try:
            plan = _asset_import_plan_serializer().loads(
                token,
                max_age=ASSET_IMPORT_PLAN_MAX_AGE_SECONDS,
            )
        except SignatureExpired as error:
            raise ValueError('The asset import preview has expired; refresh it before importing') from error
        except BadSignature as error:
            raise ValueError('The asset import preview is invalid; refresh it before importing') from error
        if plan.get('company') != current_company_code() or plan.get('user') != str(session.get('user') or ''):
            raise ValueError('The asset import preview belongs to a different user or company')
        if plan.get('rowsDigest') != _asset_import_plan_rows_digest(rows):
            raise ValueError('The import was edited after its Asset IDs were previewed; refresh the preview')
        return plan


    def _create_asset_import_row(
        payload,
        audit_timestamp,
        audit_user,
        *,
        created_asset_ids,
        inventory_target,
    ):
        created_asset_ids = list(created_asset_ids)
        expected_id_count = 1 if payload['isBulk'] else payload['quantity']
        if len(created_asset_ids) != expected_id_count:
            raise ValueError('The planned Asset IDs no longer match this import row')

        serials = (payload['serials'] + [''] * payload['quantity'])[:payload['quantity']]
        secondary_serials = (payload['secondarySerials'] + [''] * payload['quantity'])[:payload['quantity']]
        for index, asset_id in enumerate(created_asset_ids):
            asset = InventoryItem(
                asset_id=asset_id,
                brand=payload['brand'],
                model_number=payload['model'],
                version=payload['version'],
                serial_number='' if payload['isBulk'] else serials[index],
                secondary_serial_number='' if payload['isBulk'] else secondary_serials[index],
                description=payload['description'],
                is_missing=False,
                is_ooc=False,
                is_untagged=False,
                is_degraded=False,
                is_disposed=False,
                maintenance_logs=[],
                department_code=payload['department'],
                default_location=payload['defaultLocation'],
                current_location='',
                is_bulk=payload['isBulk'],
                quantity=payload['quantity'] if payload['isBulk'] else 1,
                date_of_purchase=payload['dateOfPurchase'],
                notes=payload['notes'],
                tags=payload['tags'],
            )
            mark_asset_created(asset, timestamp=audit_timestamp, user=audit_user)
            inventory_target[asset_id] = asset
        return created_asset_ids


    @app.route('/api/assets/import-template', methods=['GET'])
    @require_admin
    def download_asset_import_template():
        if not os.path.isfile(ASSET_IMPORT_TEMPLATE_PATH):
            return jsonify({'error': 'Asset import template is unavailable'}), 404
        return send_file(
            ASSET_IMPORT_TEMPLATE_PATH,
            as_attachment=True,
            download_name='Asset Import Template.csv',
            mimetype='text/csv; charset=utf-8',
        )


    @app.route('/api/assets/import-preview', methods=['POST'])
    @require_admin
    def preview_asset_import():
        uploaded = request.files.get('file')
        if not uploaded or not uploaded.filename:
            return jsonify({'error': 'Choose a CSV or Excel file to upload'}), 400
        extension = os.path.splitext(uploaded.filename)[1].lower()
        if extension not in ASSET_IMPORT_FILE_EXTENSIONS:
            return jsonify({'error': 'Upload a CSV or Excel file (.csv, .xlsx, .xlsm, or .xls)'}), 400
        file_bytes = uploaded.read(ASSET_IMPORT_MAX_BYTES + 1)
        if len(file_bytes) > ASSET_IMPORT_MAX_BYTES:
            return jsonify({'error': 'The import file must be 5 MB or smaller'}), 413
        try:
            parsed = _parse_asset_import_file(file_bytes, uploaded.filename)
            parsed.update(_asset_import_plan_response(parsed['rows']))
            existing_rejected_rows = {
                safe_int(item.get('sourceRow'), 0)
                for item in parsed['rejected']
            }
            parsed['rejected'].extend(
                {
                    'sourceRow': row.get('sourceRow', 0),
                    'error': row['departmentIssue'],
                }
                for row in parsed['rows']
                if row.get('departmentIssue')
                and safe_int(row.get('sourceRow'), 0) not in existing_rejected_rows
            )
            departments = load_departments()
            parsed['departments'] = [
                department_payload(code, departments)
                for code in sorted(departments.keys())
            ]
            return jsonify({'success': True, 'data': parsed})
        except (RuntimeError, ValueError) as error:
            return jsonify({'error': str(error)}), 400
        except Exception as error:
            logger.error('Error previewing asset import: %s', error, exc_info=True)
            return jsonify({'error': 'Could not read the asset import file'}), 500


    @app.route('/api/assets/import-plan', methods=['POST'])
    @require_admin
    def plan_asset_import():
        data = request.get_json() or {}
        rows = data.get('rows')
        if not isinstance(rows, list) or not rows:
            return jsonify({'error': 'The import does not contain any assets'}), 400
        if len(rows) > ASSET_IMPORT_MAX_ROWS:
            return jsonify({'error': f'An import can contain at most {ASSET_IMPORT_MAX_ROWS} rows'}), 400
        try:
            planned = _asset_import_plan_response(rows)
        except ValueError as error:
            return jsonify({'error': str(error)}), 400
        return jsonify({'success': True, 'data': planned})


    @app.route('/api/assets/import', methods=['POST'])
    @require_admin
    def import_assets_from_workbook():
        data = request.get_json() or {}
        raw_rows = data.get('rows')
        plan_token = str(data.get('planToken') or '')
        if not isinstance(raw_rows, list) or not raw_rows:
            return jsonify({'error': 'The import does not contain any assets'}), 400
        if len(raw_rows) > ASSET_IMPORT_MAX_ROWS:
            return jsonify({'error': f'An import can contain at most {ASSET_IMPORT_MAX_ROWS} rows'}), 400

        with inventory_action_lock:
            try:
                plan = _verify_asset_import_plan(raw_rows, plan_token)
            except ValueError as error:
                return jsonify({'error': str(error), 'code': 'asset_import_plan_required'}), 409

            inventory_changed = (
                plan.get('revision') != _asset_import_inventory_revision()
            )
            departments_changed = (
                plan.get('departmentRevision')
                != _asset_import_department_revision()
            )
            if inventory_changed or departments_changed:
                refreshed = _asset_import_plan_response(raw_rows)
                return jsonify({
                    'error': (
                        'Inventory or departments changed after this import was '
                        'previewed. Review the refreshed matches and Asset IDs, '
                        'then confirm again.'
                    ),
                    'code': 'asset_import_plan_stale',
                    'data': refreshed,
                }), 409

            normalized_rows = []
            row_errors = []
            departments = load_departments()
            try:
                new_departments = _asset_import_proposed_departments(raw_rows, departments)
            except ValueError as error:
                return jsonify({'error': str(error)}), 400
            planning_departments = {**departments, **new_departments}
            for index, raw_row in enumerate(raw_rows):
                try:
                    normalized_rows.append(_normalise_asset_import_payload(
                        raw_row,
                        departments=planning_departments,
                    ))
                except ValueError as error:
                    row_errors.append({
                        'index': index,
                        'sourceRow': safe_int((raw_row or {}).get('sourceRow'), 0) if isinstance(raw_row, dict) else 0,
                        'error': str(error),
                    })
            if row_errors:
                return jsonify({'error': 'Some imported assets need attention', 'rowErrors': row_errors}), 400

            expanded_record_count = sum(
                1 if row['isBulk'] else row['quantity']
                for row in normalized_rows
            )
            if expanded_record_count > ASSET_IMPORT_MAX_RECORDS:
                return jsonify({
                    'error': (
                        f'This import would create {expanded_record_count:,} inventory records. '
                        f'The safe limit is {ASSET_IMPORT_MAX_RECORDS:,} records per import.'
                    ),
                    'code': 'asset_import_too_large',
                }), 413

            departments_before = copy.deepcopy(departments)
            audit_timestamp = asset_audit_timestamp()
            audit_user = asset_audit_user()
            results = []
            staged_inventory = {}
            try:
                for index, payload in enumerate(normalized_rows):
                    planned_ids = list((raw_rows[index] or {}).get('assetIdsPreview') or [])
                    if any(
                        asset_id in data_manager.inventory or asset_id in staged_inventory
                        for asset_id in planned_ids
                    ):
                        raise ValueError(
                            'A planned Asset ID is no longer available; refresh the import preview'
                        )
                    created_ids = _create_asset_import_row(
                        payload,
                        audit_timestamp,
                        audit_user,
                        created_asset_ids=planned_ids,
                        inventory_target=staged_inventory,
                    )
                    results.append({
                        'index': index,
                        'sourceRow': payload.get('sourceRow', 0),
                        'assetIds': created_ids,
                        'quantity': payload['quantity'],
                        'isBulk': payload['isBulk'],
                    })
                if new_departments:
                    departments.update(new_departments)
                    save_departments(departments)
                data_manager.inventory.update(staged_inventory)
                data_manager.save_inventory()
            except Exception as error:
                for asset_id in staged_inventory:
                    data_manager.inventory.pop(asset_id, None)
                if new_departments:
                    try:
                        save_departments(departments_before)
                    except Exception:
                        logger.error(
                            'Failed to roll back departments after an asset import error',
                            exc_info=True,
                        )
                if isinstance(error, ValueError):
                    return jsonify({'error': str(error)}), 400
                logger.error('Error importing assets: %s', error, exc_info=True)
                return jsonify({'error': 'Failed to save the imported assets'}), 500

            mark_data_snapshot_current()

        if new_departments:
            mark_realtime_change('departments', {
                'action': 'created-from-asset-import',
                'departments': list(new_departments.values()),
            })

        invalidate_cache()
        created_record_count = sum(len(result['assetIds']) for result in results)
        bulk_unit_count = sum(
            result['quantity'] for result in results if result['isBulk']
        )
        log_action(
            f'Imported {len(results)} asset row(s) from a template file: '
            f'{created_record_count} inventory record(s), {bulk_unit_count} bulk unit(s)',
            system_log_only=True,
        )
        return jsonify({
            'success': True,
            'message': f'Imported {len(results)} asset row(s)',
            'data': {
                'rows': results,
                'inventoryRecordsCreated': created_record_count,
                'bulkUnitsCreated': bulk_unit_count,
            },
        })

"""Accountant-friendly Excel exports for reports and retained close packs."""
from __future__ import annotations

import io
import math

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _safe(value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub('', value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def _title(value):
    clean = ILLEGAL_CHARACTERS_RE.sub('', value)
    return ''.join(c if c not in '[]:*?/\\' else '-' for c in clean)[:31] or 'Report'


def add_report(book, report, title=None):
    sheet = book.create_sheet(_title(title or report['title']))
    sheet.sheet_view.showGridLines = False
    business = report.get('businessName') or 'Company books'
    sheet.append([_safe(business + '\n' + report['title'])])
    identity = f" · UEN {report['uen']}" if report.get('uen') else ''
    sheet.append([_safe(f"{report['period']['from']} to {report['period']['to']} · {report.get('framework', '')} · {report.get('currency', 'SGD')}{identity}")])
    sheet.append([])
    sheet.append([c['label'] for c in report['columns']])
    navy, pale, line = '18324A', 'EAF2F5', Side(style='thin', color='CCD7DD')
    sheet['A1'].font = Font(bold=True, size=15, color=navy)
    sheet['A1'].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 42
    sheet['A2'].font = Font(italic=True, color='526675')
    sheet['A2'].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[2].height = 30
    if len(report['columns']) > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(report['columns']))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(report['columns']))
    for cell in sheet[4]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.alignment = Alignment(wrap_text=True)
    for row in report['rows']:
        sheet.append([_safe(row.get(c['key'])) for c in report['columns']])
        cells = sheet[sheet.max_row]
        if row.get('section'):
            for cell in cells:
                cell.font = Font(bold=True, color=navy)
                cell.fill = PatternFill('solid', fgColor=pale)
        elif row.get('total'):
            for cell in cells:
                cell.font = Font(bold=True)
                cell.border = Border(top=line)
        for index, column in enumerate(report['columns']):
            if column.get('format') == 'money' and isinstance(cells[index].value, (int, float)):
                cells[index].number_format = '#,##0.00;[Red](#,##0.00);-'
            elif column.get('format') == 'number' and isinstance(cells[index].value, (int, float)):
                cells[index].number_format = '#,##0.########'
    notes_start = sheet.max_row + 2
    if report.get('notes'):
        sheet.append([])
        for note in report['notes']:
            sheet.append([_safe(note)])
            sheet.cell(sheet.max_row, 1).alignment = Alignment(wrap_text=True)
            if len(report['columns']) > 1:
                sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row,
                                  end_column=len(report['columns']))
    sheet.freeze_panes = 'A5'
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(report['columns']))}{max(4, 4 + len(report['rows']))}"
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:4'
    for index, column in enumerate(report['columns'], 1):
        values = [str(sheet.cell(row, index).value or '') for row in range(4, min(4 + len(report['rows']), 250) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(12, max(map(len, values), default=0) + 2), 45)
    note_width = sum(sheet.column_dimensions[get_column_letter(i)].width for i in range(1, len(report['columns']) + 1))
    for row in range(notes_start, sheet.max_row + 1):
        sheet.row_dimensions[row].height = max(30, 15 * math.ceil(len(str(sheet.cell(row, 1).value or '')) / max(1, note_width * .85)))
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:{get_column_letter(len(report["columns"]))}{sheet.max_row}'
    return sheet


def report_xlsx(report):
    book = Workbook()
    book.remove(book.active)
    add_report(book, report)
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def pack_xlsx(pack):
    book = Workbook()
    cover = book.active
    cover.title = 'Index'
    cover.sheet_view.showGridLines = False
    cover.append([pack.get('businessName') or 'Company books'])
    cover.append([pack['title']])
    cover.append([f"Period {pack['period']['from']} to {pack['period']['to']}"])
    cover.append([f"Created {pack['createdAt']} by {pack['createdBy']}"])
    cover.append([f"UEN {pack.get('uen') or 'not recorded'}"])
    cover.append([])
    cover.append(['Period review', 'Status', 'Detail'])
    for check in pack.get('checks', []):
        cover.append([check['label'], check['status'], check['detail']])
    cover.column_dimensions['A'].width = 42
    cover.column_dimensions['B'].width = 15
    cover.column_dimensions['C'].width = 90
    cover['A1'].font = Font(bold=True, size=16, color='18324A')
    cover['A2'].font = Font(bold=True, size=14)
    for report in pack['reports']:
        name = report['title']
        sheet = add_report(book, report, name)
        cover.append([name, '', 'Open report'])
        cover.cell(cover.max_row, 3).hyperlink = "#'" + sheet.title.replace("'", "''") + "'!A1"
    for row in cover:
        for cell in row:
            cell.value = _safe(cell.value)
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream

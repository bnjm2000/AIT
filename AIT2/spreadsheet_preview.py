"""Bounded, read-only worksheet previews for untrusted invoice uploads."""
from datetime import date, datetime, time
from pathlib import Path
import re
import zipfile

MAX_SHEETS = 10
MAX_ROWS = 200
MAX_COLUMNS = 30
MAX_EXPANDED_BYTES = 40 * 1024 * 1024


class PreviewLimitError(ValueError):
    """Safe, user-facing preview limit message."""


def display_value(value, number_format='General'):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d %b %Y %H:%M') if value.time() != time() else value.strftime('%d %b %Y')
    if isinstance(value, date):
        return value.strftime('%d %b %Y')
    if isinstance(value, time):
        return value.strftime('%H:%M:%S')
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        fmt = str(number_format or 'General').split(';')[0]
        if '%' in fmt:
            return f'{value * 100:,.2f}%'
        if '$' in fmt or 'SGD' in fmt.upper():
            return f'${value:,.2f}'
        decimal = re.search(r'\.([0#]+)', fmt)
        if decimal:
            text = format(value, f',.{min(len(decimal[1]), 10)}f')
            for placeholder in reversed(decimal[1]):
                if placeholder != '#' or not text.endswith('0'):
                    break
                text = text[:-1]
            return text.rstrip('.')
        if re.fullmatch(r'0{2,}', fmt):
            return str(int(value)).zfill(min(len(fmt), 30))
        if ',' in fmt:
            return format(value, ',.0f')
        return str(value) if isinstance(value, int) else format(value, '.15g')
    return str(value)[:2000]


def trim_rows(rows):
    while rows and not any(cell['text'] for cell in rows[-1]):
        rows.pop()
    used = max((index + 1 for row in rows for index, cell in enumerate(row) if cell['text']), default=0)
    return [row[:used] for row in rows]


def read_xlsx(path):
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 2000 or sum(row.file_size for row in members) > MAX_EXPANDED_BYTES:
            raise PreviewLimitError('This workbook is too large to preview safely. Download it to view it in Excel.')
        for member in members:
            if member.filename.lower().endswith(('.xml', '.rels')):
                xml = archive.read(member).replace(b'\x00', b'')
                if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b', xml, re.I):
                    raise PreviewLimitError('This workbook contains XML declarations that cannot be safely previewed.')
        for member in members:
            if member.filename.lower().endswith(('.xml', '.rels')):
                xml = archive.read(member).replace(b'\x00', b'')
                if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b', xml, re.I):
                    raise PreviewLimitError('This workbook contains XML declarations that cannot be safely previewed.')
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        visible = [sheet for sheet in formulas if sheet.sheet_state == 'visible']
        sheets = []
        for sheet in visible[:MAX_SHEETS]:
            row_count = min(sheet.max_row or MAX_ROWS, MAX_ROWS)
            column_count = min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS)
            rows = []
            values = cached[sheet.title].iter_rows(max_row=row_count, max_col=column_count)
            for formula_row, value_row in zip(sheet.iter_rows(max_row=row_count, max_col=column_count), values):
                row = []
                for cell, saved in zip(formula_row, value_row):
                    is_formula = getattr(cell, 'data_type', '') == 'f'
                    missing = is_formula and saved.value is None
                    value = saved.value if is_formula else cell.value
                    row.append({
                        'text': 'Not calculated' if missing else display_value(value, getattr(cell, 'number_format', 'General')),
                        'bold': bool(getattr(getattr(cell, 'font', None), 'bold', False)),
                        'numeric': isinstance(value, (int, float)) and not isinstance(value, bool),
                        'missing': missing,
                    })
                rows.append(row)
            sheets.append({'name': sheet.title, 'rows': trim_rows(rows),
                           'truncated': (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS})
        return {'sheets': sheets, 'truncatedSheets': len(visible) > MAX_SHEETS}
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def read_xls(path):
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
    try:
        visible = [index for index in range(workbook.nsheets) if workbook.sheet_by_index(index).visibility == 0]
        sheets = []
        for index in visible[:MAX_SHEETS]:
            sheet = workbook.sheet_by_index(index)
            rows = []
            for row_index in range(min(sheet.nrows, MAX_ROWS)):
                row = []
                for column_index in range(min(sheet.ncols, MAX_COLUMNS)):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#ERROR')
                    xf = workbook.xf_list[cell.xf_index]
                    fmt = workbook.format_map.get(xf.format_key)
                    row.append({'text': display_value(value, fmt.format_str if fmt else 'General'),
                                'bold': bool(workbook.font_list[xf.font_index].bold),
                                'numeric': isinstance(value, (int, float)) and not isinstance(value, bool),
                                'missing': False})
                rows.append(row)
            sheets.append({'name': sheet.name, 'rows': trim_rows(rows),
                           'truncated': sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS})
            workbook.unload_sheet(index)
        return {'sheets': sheets, 'truncatedSheets': len(visible) > MAX_SHEETS}
    finally:
        workbook.release_resources()


def read_spreadsheet_preview(path):
    if Path(path).stat().st_size > 10 * 1024 * 1024:
        raise ValueError('This workbook is too large to preview. Download it to view it in Excel.')
    try:
        if Path(path).suffix.lower() == '.xlsx':
            return read_xlsx(path)
        if Path(path).suffix.lower() == '.xls':
            return read_xls(path)
        raise PreviewLimitError('Only .xlsx and .xls workbooks can be previewed.')
    except PreviewLimitError:
        raise
    except Exception as exc:
        raise ValueError('This workbook could not be previewed. It may be damaged or password-protected. Download the original to open it in Excel.') from exc

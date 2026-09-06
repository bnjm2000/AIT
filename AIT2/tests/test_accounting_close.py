import io
import unittest

from openpyxl import load_workbook
from tests import test_accounting_workspace as workspace_tests
from services.accounting_close import REVIEW_ITEMS


class AccountingCloseTests(unittest.TestCase):
    setUp = workspace_tests.AccountingWorkspaceTests.setUp
    tearDown = workspace_tests.AccountingWorkspaceTests.tearDown
    login = workspace_tests.AccountingWorkspaceTests.login
    make_user = workspace_tests.AccountingWorkspaceTests.make_user
    request = workspace_tests.AccountingWorkspaceTests.request
    begin = workspace_tests.AccountingWorkspaceTests.begin
    document = workspace_tests.AccountingWorkspaceTests.document
    post_document = workspace_tests.AccountingWorkspaceTests.post_document
    workspace = workspace_tests.AccountingWorkspaceTests.workspace
    report = workspace_tests.AccountingWorkspaceTests.report

    def close(self, query='from=2026-01-01&to=2026-03-31'):
        response = self.client.get('/api/finance/accounting/close?' + query)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.get_json()['close']

    def filing(self, **extra):
        value = {'type': 'gst', 'from': '2026-01-01', 'to': '2026-03-31', 'dueDate': '2026-04-30',
                 'assignedTo': 'bnjm2000', 'checklist': {key: True for key in REVIEW_ITEMS},
                 'declarations': {str(key): 0 for key in (9, 10, 11, 12, 14, 15, 16, 17)},
                 'reviewNotes': 'Reviewed outstanding receipts, purchase support and unposted drafts with the finance team.'}
        value.update(extra)
        return self.request('actions/filing-save', value)['record']

    def gst_on(self):
        self.begin()
        self.request('settings', {'gstRegistered': True}, method='put')

    def test_gst_working_paper_adjustments_declarations_and_no_ledger_posting(self):
        self.gst_on()
        self.post_document(self.document(amount=1000, lines=[dict(description='Sale', quantity=1, unitPrice=1000, taxCode='SR9')]))
        count = len(self.workspace()['journals'])
        draft = self.filing(adjustments={'1': 50, '6': 4.5}, adjustmentReason='Deemed supply per schedule GST-03')
        rows = {r['box']: r for r in draft['gstRows']}
        self.assertEqual(rows[1]['ledger'], 1000)
        self.assertEqual(rows[1]['amount'], 1050)
        self.assertEqual(rows[4]['amount'], 1050)
        self.assertEqual(rows[8]['amount'], 94.5)
        self.assertEqual(len(self.workspace()['journals']), count)
        reviewed = self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']))['record']
        self.assertEqual(reviewed['status'], 'reviewed')
        self.assertEqual(self.report('gst')['rows'][0]['amount'], 1000)

    def test_gst_review_requires_explicit_special_declarations_and_complete_checks(self):
        self.gst_on()
        draft = self.filing(declarations={})
        self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']), expected=400)
        draft = self.filing(id=draft['id'], version=draft['version'], checklist={})
        self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']), expected=400)
        self.request('actions/filing-save', dict(type='gst', **{'from':'2026-01-01','to':'2026-03-31'},
                     dueDate='2026-04-30', adjustments={'7':20}), expected=400)

    def test_filing_rejects_changed_books_stale_edits_and_independent_self_review(self):
        self.gst_on()
        self.request('actions/controls', dict(approvalRequired=True))
        self.login('sales-admin')
        draft = self.filing()
        self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']), expected=400)
        self.begin()
        reviewed = self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']))['record']
        self.request('actions/filing-save', dict(type='gst', id=draft['id'], version=draft['version']), expected=400)
        self.document()
        self.assertTrue(self.close()['filings'][0]['booksChanged'])
        refreshed = self.filing(id=draft['id'], version=reviewed['version'])
        self.request('actions/filing-review', dict(id=refreshed['id'], version=refreshed['version']), expected=400)
        self.login('review-admin')
        self.request('actions/filing-review', dict(id=refreshed['id'], version=refreshed['version']))

    def test_external_filing_requires_review_acknowledgement_and_retains_history(self):
        self.gst_on()
        draft = self.filing()
        self.request('actions/filing-file', dict(id=draft['id'], version=draft['version']), expected=400)
        reviewed = self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']))['record']
        self.request('actions/filing-file', dict(id=draft['id'], version=reviewed['version'], filedDate='2026-04-20', reference='ACK-123'), expected=400)
        upload = self.client.post(f'/api/finance/accounting/attachments/filings/{draft["id"]}', data={'file': (io.BytesIO(b'Portal acknowledgement TEST'), 'ack.txt')})
        self.assertEqual(upload.status_code, 200)
        attachment = upload.get_json()['record']
        package = self.client.post(f'/api/finance/accounting/attachments/filings/{draft["id"]}',
                                  data={'file': (io.BytesIO(b'<?xml version="1.0"?><xbrl/>'), 'financial-statements.xbrl')})
        self.assertEqual(package.status_code, 200)
        filed = self.request('actions/filing-file', dict(id=draft['id'], version=reviewed['version'], filedDate='2026-04-20', reference='ACK-123', attachmentId=attachment['id']))['record']
        self.assertEqual(filed['status'], 'filed')
        self.request('actions/filing-save', dict(id=filed['id'], version=filed['version']), expected=400)
        self.post_document(self.document())
        stored = self.close()['filings'][0]
        self.assertTrue(stored['booksChanged'])
        self.assertEqual(stored['externalReference'], 'ACK-123')
        self.assertEqual(stored['gstRows'], filed['gstRows'])
        correction = self.filing(type='gst-correction', relatedFilingId=filed['id'])
        self.assertEqual(correction['relatedFilingId'], filed['id'])

    def test_filing_cannot_use_another_records_acknowledgement(self):
        self.gst_on()
        one, two = self.filing(), self.filing()
        reviewed = self.request('actions/filing-review', dict(id=one['id'], version=one['version']))['record']
        upload = self.client.post(f'/api/finance/accounting/attachments/filings/{two["id"]}', data={'file': (io.BytesIO(b'Unrelated'), 'ack.txt')}).get_json()['record']
        self.request('actions/filing-file', dict(id=one['id'], version=reviewed['version'], filedDate='2026-04-20', reference='WRONG', attachmentId=upload['id']), expected=400)

    def test_report_pack_is_retained_and_excel_contains_the_original_totals(self):
        self.begin()
        self.post_document(self.document(amount=100))
        pack = self.request('actions/pack-create', dict(title='=UNTRUSTED()', **{'from':'2026-01-01','to':'2026-03-31'}))['record']
        self.post_document(self.document(amount=200))
        path = f'/api/finance/accounting/packs/{pack["id"]}'
        saved = self.client.get(path+'/json').get_json()
        pnl = next(r for r in saved['reports'] if r['name']=='profit-loss')
        self.assertEqual(next(r['amount'] for r in pnl['rows'] if r.get('key')=='net-profit'), 100)
        self.assertEqual(len(saved['sourceSnapshot']['documents']), 1)
        response = self.client.get(path+'/xlsx')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data), data_only=False)
        self.assertEqual(book['Index']['A2'].data_type, 's')
        self.assertIn('Profit & Loss', book.sheetnames)
        self.assertEqual(book['Profit & Loss'].freeze_panes, 'A5')

    def test_close_checks_detect_control_difference_drafts_and_missing_purchase_evidence(self):
        self.begin()
        self.request('journals', dict(status='posted', date='2026-01-10', lines=[dict(accountCode='1100', debit=50), dict(accountCode='4000', credit=50)]), expected=201)
        self.post_document(self.document('bill'))
        self.document()
        checks = {c['id']: c for c in self.close()['checks']}
        self.assertEqual(checks['ar']['count'], 50)
        self.assertEqual(checks['drafts']['count'], 1)
        self.assertEqual(checks['evidence']['count'], 1)

    def test_bank_reconciliation_includes_earlier_unmatched_items_and_outstanding_cash(self):
        self.begin()
        self.post_document(self.document('receive_money', amount=100))
        self.request('bank-transactions', dict(date='2026-01-15', amount=70, bankAccount='1000', description='Unbooked receipt'), expected=201)
        self.request('actions/bank-closing', dict(date='2026-03-31', bankAccount='1000', closingBalance=70, reference='Statement March'))
        report = self.client.get('/api/finance/accounting/reports/bank-reconciliation?from=2026-03-01&to=2026-03-31').get_json()['report']
        self.assertEqual(report['reconciliation']['difference'], 0)
        self.assertEqual(report['reconciliation']['unmatchedStatementCount'], 1)
        self.assertTrue(any('Unbooked receipt' in r['label'] for r in report['rows']))
        self.assertEqual(next(c for c in self.close()['checks'] if c['id']=='bank-1000')['status'], 'review')

    def test_supplier_statement_currency_and_fiscal_year_columns(self):
        self.begin()
        self.post_document(self.document('bill', currency='USD', exchangeRate=1.3, amount=100))
        report = self.report('supplier-statement', '&contact=Example%20Pte%20Ltd&currency=USD')
        self.assertEqual(report['currency'], 'USD')
        self.assertEqual(report['rows'][-1]['balance'], 100)
        self.request('settings', dict(financialYearStartMonth=4), method='put')
        annual = self.report('profit-loss', '&groupBy=yearly')
        self.assertEqual([c['label'] for c in annual['columns']][1:], ['FY to 2026-03-31', 'FY to 2026-12-31'])

    def test_operating_revenue_mapping_excludes_other_income_from_gst_box13(self):
        self.gst_on()
        self.post_document(self.document(amount=100))
        self.post_document(self.document(amount=30, lines=[dict(description='Grant', quantity=1, unitPrice=30, accountCode='4100')]))
        self.assertEqual(self.report('gst')['rows'][12]['amount'], 100)
        self.request('actions/controls', dict(gstRevenueAccounts=['4000','4100']))
        self.assertEqual(self.report('gst')['rows'][12]['amount'], 130)
        self.request('actions/controls', dict(gstRevenueAccounts=['1000']), expected=400)

    def test_tax_reconciliation_and_export_are_working_papers_not_automatic_tax(self):
        self.begin()
        self.post_document(self.document(amount=1000))
        draft = self.filing(type='income-tax', yearOfAssessment='2027', declaredIncome=1050,
                           taxAdjustments=[dict(label='Depreciation add-back', amount=100, reference='FA-1'), dict(label='Capital allowances', amount=-50, reference='CA-1')])
        self.assertEqual(draft['adjustedProfit'], 1050)
        response = self.client.get(f'/api/finance/accounting/filings/{draft["id"]}/xlsx')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data))
        self.assertTrue(any('Declared' in str(c.value) or 'declared' in str(c.value) for row in book.active for c in row))

    def test_contact_details_are_editable_but_transacted_names_stay_stable(self):
        self.begin()
        contact = self.request('actions/contacts', dict(name='Example Pte Ltd', email='old@example.test'))['record']
        self.post_document(self.document())
        self.request('actions/contacts', dict(contact, email='new@example.test'))
        self.request('actions/contacts', dict(contact, name='Changed'), expected=400)
        self.assertEqual(self.workspace()['workspace']['contacts'][0]['email'], 'new@example.test')

    def test_non_admin_cannot_read_export_or_change_filings(self):
        self.gst_on()
        self.request('actions/controls', dict(roles={'no-sales':'auditor'}))
        draft = self.filing()
        self.login('no-sales')
        self.assertEqual(self.client.get(f'/api/finance/accounting/filings/{draft["id"]}/xlsx').status_code, 403)
        self.assertEqual(self.client.get('/api/finance/accounting/close?from=2026-01-01&to=2026-03-31').status_code, 403)
        self.request('actions/filing-review', dict(id=draft['id'], version=draft['version']), expected=403)
        self.request('actions/pack-create', {}, expected=403)

    def test_tax_inclusive_and_blocked_gst_preserve_cost_and_exclude_input_claim(self):
        self.gst_on()
        bill = self.post_document(self.document('bill', priceBasis='inclusive', lines=[
            dict(description='Claimable purchase', quantity=1, unitPrice=109, taxCode='TX9'),
            dict(description='Disallowed input tax', quantity=1, unitPrice=109, taxCode='BL9')]))
        self.assertEqual(bill['total'], 218)
        self.assertEqual(bill['net'], 200)
        self.assertEqual(bill['tax'], 18)
        journal = next(j for j in self.workspace()['journals'] if j['id'] == bill['journalId'])
        self.assertEqual(sum(r['debit'] for r in journal['lines'] if r['accountCode']=='6800'), 209)
        self.assertEqual(sum(r['debit'] for r in journal['lines'] if r['accountCode']=='1200'), 9)
        gst = self.report('gst')['rows']
        self.assertEqual(gst[4]['amount'], 100)
        self.assertEqual(gst[6]['amount'], 9)

    def test_stock_receipt_cost_keeps_fractional_unit_precision(self):
        self.begin()
        self.request('actions/controls', dict(inventoryEnabled=True))
        item = self.request('actions/stock-items', dict(sku='THIRDS', name='Three pieces'))['record']
        self.post_document(self.document('bill', lines=[dict(description='Stock', quantity=3, unitPrice='33.3333333333', stockItemId=item['id'])]))
        inventory = self.report('inventory')['rows'][0]
        self.assertEqual(inventory['value'], 100)
        self.assertTrue(inventory['ref']['journalIds'])

    def test_direct_cash_can_repay_a_liability_but_cannot_bypass_ap(self):
        self.begin()
        doc = self.post_document(self.document('expense', lines=[dict(description='GST payment', quantity=1, unitPrice=100, accountCode='2100')]))
        journal = next(j for j in self.workspace()['journals'] if j['id']==doc['journalId'])
        self.assertEqual(next(r['debit'] for r in journal['lines'] if r['accountCode']=='2100'), 100)
        self.request('documents', dict(kind='expense', date='2026-01-20', contact='Vendor', lines=[dict(description='Invalid AP shortcut', quantity=1, unitPrice=100, accountCode='2000')]), expected=400)

    def test_no_filing_required_decision_needs_reason_and_is_retained(self):
        self.begin()
        draft = self.filing(type='eci', yearOfAssessment='2027')
        self.request('actions/filing-not-required', dict(id=draft['id'], version=draft['version']), expected=400)
        result = self.request('actions/filing-not-required', dict(id=draft['id'], version=draft['version'], reason='ECI waiver criteria assessed; see schedule ECI-1'))['record']
        self.assertEqual(result['status'], 'not-required')
        self.request('actions/filing-save', dict(id=draft['id'], version=result['version']), expected=400)

    def test_direct_cash_correction_preserves_historical_gst_and_reverses_current_balance(self):
        self.gst_on()
        doc = self.post_document(self.document('expense', lines=[dict(description='Incorrect cash entry', quantity=1, unitPrice=100, taxCode='TX9')]))
        self.request('actions/reverse-cash', dict(id=doc['id'], date='2026-02-01', reason='Wrong cash account selected'))
        self.request('actions/reverse-cash', dict(id=doc['id'], date='2026-02-01', reason='Duplicate correction'), expected=400)
        self.assertEqual(self.report('gst')['rows'][6]['amount'], 0)
        january = self.client.get('/api/finance/accounting/reports/gst?from=2026-01-01&to=2026-01-31').get_json()['report']
        self.assertEqual(january['rows'][6]['amount'], 9)
        self.assertEqual(self.workspace()['summary']['cash'], 0)

    def test_variance_drilldown_includes_current_and_comparison_journals(self):
        self.begin()
        previous = self.post_document(self.document(amount=50, date='2025-01-15'))
        current = self.post_document(self.document(amount=100))
        report = self.report('profit-loss', '&compare=year')
        row = next(r for r in report['rows'] if r.get('key')=='4000')
        self.assertEqual(row['variance'], 50)
        self.assertEqual(set(row['cellRefs']['variance']['journalIds']), {current['journalId'], previous['journalId']})


if __name__ == '__main__':
    unittest.main()
